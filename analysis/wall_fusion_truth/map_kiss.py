#!/usr/bin/env python3
"""Feasibility stage of the Kiss et al. 2019 truth set (plan Task 5.1; spec D8, D9, D20).

No model is run. The script answers one question: how much of the Kiss 2019 hypha-gene
table can be turned into protein sequences, and how many clades could be held out.

Steps (all deterministic; every network answer is cached in results/query_cache.json):
 1. Resolve the 16 JGI organism codes of the table to NCBI taxa (Kiss Supplementary
    Data 1 gives the species name; NCBI Taxonomy gives the taxon ID and the lineage).
 2. Map each row to a UniProt sequence. There is NO route by JGI protein ID (JGI pages
    are behind a challenge page, UniProt has no JGI cross-reference, the HPCC proteomes
    use their own IDs). A row is therefore mapped only by an identifier that the row itself
    carries (a locus tag such as Afu3g12700, NCU09332, BATDEDRAFT_20500) inside the genome
    strain of the paper. A unique gene-name hit is kept apart as a CANDIDATE and is never
    counted as mapped.
 3. Optional (--jgi-dir): look each JGI protein ID up in the gene-catalog protein FASTA file
    of its MycoCosm portal (files that the owner downloaded; they are NOT in the repository).
    A row is mapped by file (status mapped_jgi_file) only when the header name of the protein
    agrees with the row, with the UniProt candidate, or (for rows that carry no name evidence)
    when the portal as a whole is validated by its other rows. Only the sequence length and
    sha256 are written, never the sequence.
 4. Count rows, mapping loss, clades and clusters per slice, and the clades that a floor of
    3, 5 or 8 independent clusters admits for hold-out.
 5. Count UniProt reviewed entries of the wall and fusion Pfam families in Basidiomycota
    and other clades (the second source).

Independent unit: the Kiss `cluster_number` is a PROXY. It is a family from orthology and
MCL clustering, not a 30 percent identity cluster. The MMseqs2 clustering of D8 is Task 5.3.

Untrusted input: the Kiss supplement and all network answers are data only. Nothing from
them is executed. The optional xlsx step needs openpyxl and is run only by hand.

Usage:
    python3 map_kiss.py                 # offline, from the committed cache
    python3 map_kiss.py --online        # fill the cache from public NCBI and UniProt APIs
    python3 map_kiss.py --jgi-dir to_import/kiss2019_jgi   # add the JGI-file step
    python3 map_kiss.py --species-xlsx MOESM5.xlsx   # rebuild results/kiss_species_72.tsv
"""

from __future__ import annotations

import argparse
import csv
import difflib
import gzip
import hashlib
import json
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
TSV_DEFAULT = ROOT / "docs/reports/cell-wall-fusion/06-kiss2019-hypha-genes.tsv"
MODELS_DEFAULT = ROOT / "analysis/wall_fusion_report/wall_fusion_models.tsv"
OUT_DEFAULT = HERE / "results"

FLOORS = (3, 5, 8)
# Owner decisions of 2026-10-22: the fine clade level is the primary hold-out and the coarse
# level a second, harder test; a stratum of clusters is `estimated` only with at least
# ESTIMATED_FLOOR clusters of its own (fewer: `descriptive`); the fusion slice is always
# descriptive; a family representative (no clade) is never a test control.
PRIMARY_LEVEL = "fine"
ESTIMATED_FLOOR = 5
DESCRIPTIVE_SLICES = ("fusion",)
NCBI = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/"
UNIPROT = "https://rest.uniprot.org/uniprotkb/search"
UNIPROT_FIELDS = "accession,reviewed,id,gene_names,organism_name,organism_id,length,sequence"

# Genome strains named in Kiss Supplementary Data 1 ("Genome version" column). A strain is
# used only where the table text names it; the others use the species taxon.
STRAIN_TERMS = {
    "Aspfu1": "Aspergillus fumigatus Af293",
    "Neucr2": "Neurospora crassa OR74A",
    "Canalb1": "Candida albicans SC5314",
    "Sacce1": "Saccharomyces cerevisiae S288C",
    "Ustma22": "Ustilago maydis 521",
    "Hisca1": "Histoplasma capsulatum NAm1",
    "Ashgo11": "Eremothecium gossypii ATCC 10895",
    "Schco3": "Schizophyllum commune H4-8",
    "CryneH991": "Cryptococcus neoformans var. grubii H99",
    "Batde5": "Batrachochytrium dendrobatidis JAM81",
    "Crypa2": "Cryphonectria parasitica EP155",
}

SLICE_OF_CATEGORY = {
    "adhesion": "adhesion",
    "cell wall": "wall",
    "cell wall (GPI anchor assembly)": "wall",
    "cell fusion": "fusion",
    "cell fusion (syncytium formation)": "fusion",
    "pheromone receptor (mating, fusion-adjacent)": "fusion",
}
SLICES = ("adhesion", "wall", "fusion")

# Locus tag patterns that a row can carry in `gene_or_family_id`. The Afu form is turned
# into the UniProt form AFUA_<n>G<nnnnn> by one fixed rule.
LOCUS_PATTERNS = (
    re.compile(r"(?:^|_|/)(Afu\d+g\d+)(?:$|_|\s)"),
    re.compile(r"^(NCU\d{5})$"),
    re.compile(r"^(AGOS_[A-Z0-9]+)$"),
    re.compile(r"^(BATDEDRAFT_\d+)$"),
    re.compile(r"^(CN[A-Z]\d{5})$"),
    re.compile(r"^(UM\d{5})$"),
    re.compile(r"^(ACR\d{3}[WC])$"),
)
# Kiss organism code -> portal directory and header name of the downloaded MycoCosm file.
JGI_PORTAL = {
    "Ustma22": "Ustma2_2",
    "CryneH991": "Cryne_H99_1",
    "Ashgo11": "Ashgo1_1",
    "Botci1": "Bcine1",
    "Copci": "Copci_AmutBmut1",
}
# A portal validates its rows without name evidence when at least PORTAL_MIN_AGREE rows agree
# and at most PORTAL_MAX_DIFFER_RATE of the rows with evidence disagree.
PORTAL_MIN_AGREE, PORTAL_MAX_DIFFER_RATE = 3, 0.10
SIMILAR_RATIO = 0.90  # difflib ratio from which two different sequences count as 'similar'
SIMILAR_MAX_LEN = 4000
JGI_COLS = [
    "map_status_uniprot",
    "jgi_portal",
    "jgi_status",
    "jgi_name",
    "jgi_length",
    "jgi_seq_sha256",
    "jgi_check",
    "jgi_check_how",
    "jgi_vs_uniprot",
    "jgi_accept",
    "jgi_locus_tag",
    "jgi_locus_result",
    "jgi_locus_accession",
    "jgi_locus_genes",
]
# A header name that is a locus tag: upper-case scheme prefix, then digits; no dots, no '#'.
LOCUS_LIKE = re.compile(r"^[A-Z]{2,8}_[A-Z0-9]*\d[A-Za-z0-9]*$")
NAME_STOP = {"like"}
MIN_LEN, MAX_LEN, MAX_X_RUN = 20, 15000, 5
AA = set("ACDEFGHIKLMNPQRSTVWYXUBZO")


# ----------------------------------------------------------------------------- pure logic


def slice_of(category: str) -> str | None:
    return SLICE_OF_CATEGORY.get(category.strip())


def organism_code(jgi_id: str) -> str:
    """JGI protein ID '4407_Aspfu1' -> 'Aspfu1'. Empty string for no ID."""
    m = re.match(r"^(\d+)_(.+)$", jgi_id.strip())
    return m.group(2) if m else ""


def afu_uniprot_form(afu: str) -> str:
    """Afu3g12700 -> AFUA_3G12700 (fixed rule of the A. fumigatus Af293 locus tags)."""
    m = re.match(r"^Afu(\d+)g(\d+)$", afu)
    return f"AFUA_{m.group(1)}G{m.group(2)}" if m else afu


def locus_tags(gene_field: str) -> list[str]:
    """Locus tags that the row itself carries, in UniProt form, in order, no duplicates."""
    out: list[str] = []
    field = gene_field.strip()
    for pat in LOCUS_PATTERNS:
        for m in pat.finditer(field):
            tag = afu_uniprot_form(m.group(1)) if m.group(1).startswith("Afu") else m.group(1)
            if tag not in out:
                out.append(tag)
    return out


def gene_name_tokens(gene_field: str) -> list[str]:
    """Gene-name tokens of a row (split on '_' and '/'), without locus tags and stop words."""
    if locus_tags(gene_field):
        return []
    toks = [t for t in re.split(r"[_/]", gene_field.strip()) if t]
    return [t for t in toks if t.lower() not in NAME_STOP and len(t) >= 2]


def plausible_protein(seq: str) -> tuple[bool, str]:
    if not (MIN_LEN <= len(seq) <= MAX_LEN):
        return False, f"length {len(seq)} outside {MIN_LEN}-{MAX_LEN}"
    if not set(seq) <= AA:
        return False, "non amino-acid characters"
    if re.search(rf"X{{{MAX_X_RUN},}}", seq):
        return False, f"X run of {MAX_X_RUN} or more"
    return True, ""


def pick_entry(hits: list[dict]) -> tuple[dict | None, str]:
    """One unique entry from hits (reviewed preferred), else (None, reason)."""
    if not hits:
        return None, "no_hit"
    reviewed = [h for h in hits if h["reviewed"]]
    pool = reviewed or hits
    accs = sorted({h["accession"] for h in pool})
    if len(accs) != 1:
        return None, "ambiguous"
    return next(h for h in pool if h["accession"] == accs[0]), ""


def map_row(row: dict, scope_taxon: str, lookup) -> dict:
    """Map one row. `lookup(kind, term, taxon)` returns a list of UniProt entry dicts.

    status values:
      mapped_locus_tag      accepted: a locus tag of the row matches one UniProt entry in the strain
      mapped_family_rep     accepted, family-level row without JGI ID: A. fumigatus locus tag
                            of the ortholog column (representative of the family, not the Kiss protein)
      candidate_gene_name   NOT accepted: one UniProt entry matches the gene name
      lost                  not mapped (see reason)
    """
    jgi = row["protein_id_jgi"].strip()
    res = {
        "status": "lost",
        "reason": "",
        "accession": "",
        "match_kind": "",
        "match_term": "",
        "length": "",
        "uniprot_organism_id": "",
        "id_consistent": "",
        "sequence": "",
        "genes": "",
    }
    if not jgi:
        afus = re.findall(r"Afu\d+g\d+", row.get("ortholog_Afu", ""))
        if not afus:
            res["reason"] = "no_jgi_id_and_no_afu_locus"
            return res
        tag = afu_uniprot_form(afus[0])
        entry, why = pick_entry(lookup("gene_exact", tag, scope_taxon))
        res["match_kind"], res["match_term"] = "afu_locus_of_family", tag
        if entry is None:
            res["reason"] = f"no_jgi_id; afu locus {why}"
            return res
        return _accept(res, entry, "mapped_family_rep")

    tags = locus_tags(row["gene_or_family_id"])
    for tag in tags:
        res["match_kind"], res["match_term"] = "row_locus_tag", tag
        entry, why = pick_entry(lookup("gene_exact", tag, scope_taxon))
        if entry is not None:
            digits = re.sub(r"\D", "", tag)
            res["id_consistent"] = "yes" if jgi.split("_")[0] == digits else ""
            return _accept(res, entry, "mapped_locus_tag")
        res["reason"] = f"locus_tag_{why}_in_genome_strain"
    if tags:
        return res

    names = gene_name_tokens(row["gene_or_family_id"])
    accs: dict[str, dict] = {}
    ambiguous = False
    for tok in names:
        entry, why = pick_entry(lookup("gene_exact", tok, scope_taxon))
        ambiguous = ambiguous or why == "ambiguous"
        if entry is not None:
            accs[entry["accession"]] = entry
            res["match_term"] = tok
    res["match_kind"] = "gene_name"
    if len(accs) == 1 and not ambiguous:
        return _accept(res, next(iter(accs.values())), "candidate_gene_name")
    res["reason"] = "gene_name_ambiguous" if (ambiguous or len(accs) > 1) else "gene_name_no_hit"
    return res


def _accept(res: dict, entry: dict, status: str) -> dict:
    ok, why = plausible_protein(entry["sequence"])
    res.update(
        accession=entry["accession"],
        length=str(len(entry["sequence"])),
        uniprot_organism_id=str(entry["organism_id"]),
        genes=entry.get("genes", ""),
    )
    if not ok:
        res["status"], res["reason"] = "lost", f"implausible_sequence: {why}"
        return res
    res["status"], res["reason"], res["sequence"] = status, "", entry["sequence"]
    return res


def jgi_loss_reason(row: dict, mapped: dict) -> str:
    """The loss reason of a row that has a JGI ID but no accepted mapping."""
    return "jgi_id_has_no_public_sequence_route; " + (mapped["reason"] or mapped["status"])


def count_rows(rows: list[dict]) -> dict:
    """Row counts of the mapping stage (used by the report and the unit test)."""
    n = len(rows)
    with_id = sum(1 for r in rows if r["protein_id_jgi"].strip())
    c = defaultdict(int)
    for r in rows:
        c[r["map_status"]] += 1
    lost = [r for r in rows if r["map_status"] == "lost"]
    reasons = defaultdict(int)
    for r in lost:
        reasons[r["map_reason"]] += 1
    return {
        "rows": n,
        "with_jgi_id": with_id,
        "without_jgi_id": n - with_id,
        "mapped_locus_tag": c["mapped_locus_tag"],
        "mapped_family_rep": c["mapped_family_rep"],
        "mapped_jgi_file": c["mapped_jgi_file"],
        "candidate_gene_name": c["candidate_gene_name"],
        "distinct_accessions_locus_tag": len(
            {r["accession"] for r in rows if r["map_status"] == "mapped_locus_tag"}
        ),
        "lost": len(lost),
        "lost_reasons": dict(sorted(reasons.items())),
    }


def clade_tables(rows: list[dict], include) -> dict:
    """{(slice, clade_level, clade): {rows, clusters, shared}} for rows passing include(row)."""
    out: dict = {}
    for level in ("coarse", "fine"):
        key = f"clade_{level}"
        per = defaultdict(lambda: {"rows": 0, "clusters": set()})
        clade_by_cluster = defaultdict(set)
        for r in rows:
            if not include(r) or not r["slice"] or not r[key]:
                continue
            per[(r["slice"], r[key])]["rows"] += 1
            per[(r["slice"], r[key])]["clusters"].add(r["cluster_number"])
            clade_by_cluster[(r["slice"], r["cluster_number"])].add(r[key])
        for (sl, clade), d in per.items():
            shared = sum(1 for c in d["clusters"] if len(clade_by_cluster[(sl, c)]) > 1)
            out[(sl, level, clade)] = {
                "rows": d["rows"],
                "clusters": len(d["clusters"]),
                "shared": shared,
            }
    return out


def stratum_status(n_clusters: int, sl: str) -> str:
    """`estimated` only with ESTIMATED_FLOOR clusters of the stratum itself; fusion never; `none` for 0."""
    if n_clusters == 0:
        return "none"
    if sl in DESCRIPTIVE_SLICES or n_clusters < ESTIMATED_FLOOR:
        return "descriptive"
    return "estimated"


def level_role(level: str) -> str:
    return "primary" if level == PRIMARY_LEVEL else "second_harder_test"


def control_role(row: dict) -> str:
    """family_rep: a family representative (training seed and baseline report only, never a
    test control, no clade, never counted toward a floor); kiss_row: a row of a Kiss protein
    (whether it is mapped is in map_status); unmapped_family_row: a family row without locus."""
    if row["map_status"] == "mapped_family_rep":
        return "family_rep"
    if row["organism_code"]:
        return "kiss_row"
    return "unmapped_family_row"


def family_rep_clusters(rows: list[dict]) -> list[list]:
    """[slice, cluster, rep rows, strictly mapped Kiss rows] for clusters with a family rep."""
    strict = ("mapped_locus_tag", "mapped_jgi_file")
    out = []
    keys = sorted(
        {(r["slice"], r["cluster_number"]) for r in rows if r["control_role"] == "family_rep"},
        key=lambda k: (k[0], int(k[1])),
    )
    for sl, cl in keys:
        same = [r for r in rows if r["slice"] == sl and r["cluster_number"] == cl]
        out.append(
            [
                sl,
                cl,
                sum(1 for r in same if r["control_role"] == "family_rep"),
                sum(
                    1 for r in same if r["control_role"] == "kiss_row" and r["map_status"] in strict
                ),
            ]
        )
    return out


NAME_CONFLICT_COLS = [
    "gene_or_family_id",
    "protein_id_jgi",
    "slice",
    "jgi_header_name",
    "file_protein_accession",
    "file_protein_gene_names",
    "file_protein_length",
    "named_candidate_accession",
    "named_candidate_length",
    "why",
]


def name_conflict_table(rows: list[dict]) -> list[list]:
    """Review list: rows whose JGI protein is a UniProt entry other than the row's gene name."""
    out = []
    for r in sorted(rows, key=lambda r: (r["slice"], r["organism_code"], r["gene_or_family_id"])):
        if r["jgi_locus_result"] != "confirmed_name_conflict":
            continue
        why = (
            f"the protein at this ID is {r['jgi_locus_accession']} ({r['jgi_locus_tag']}), "
            f"not the UniProt entry {r['accession']} of the name {r['match_term']}; "
            "the JGI sequence equals the locus-tag entry"
        )
        out.append(
            [
                r["gene_or_family_id"],
                r["protein_id_jgi"],
                r["slice"],
                r["jgi_name"],
                r["jgi_locus_accession"],
                r["jgi_locus_genes"],
                r["jgi_length"],
                r["accession"],
                r["length"],
                why,
            ]
        )
    return out


def eligible(tables: dict, sl: str, level: str, floor: int) -> list[tuple[str, int, int]]:
    """Clades with at least `floor` independent clusters: (clade, clusters, rows)."""
    return sorted(
        (clade, v["clusters"], v["rows"])
        for (s, lv, clade), v in tables.items()
        if s == sl and lv == level and v["clusters"] >= floor
    )


# ----------------------------------------------------------------------------- JGI-file step


def jgi_portal(code: str) -> str:
    """Kiss organism code -> portal name of the downloaded file (identical for most codes)."""
    return JGI_PORTAL.get(code, code)


def jgi_protein_id(jgi_id: str) -> str:
    """'4407_Aspfu1' -> '4407'."""
    m = re.match(r"^(\d+)_.+$", jgi_id.strip())
    return m.group(1) if m else ""


def read_jgi_table(path: Path) -> dict[str, dict]:
    """The file table of the downloaded set: {portal: row} for the files marked 'use'."""
    used: dict[str, dict] = {}
    for r in read_tsv(path):
        if not r["role"].startswith("use"):
            continue
        if r["portal"] in used:
            raise ValueError(f"two files marked 'use' for portal {r['portal']}")
        used[r["portal"]] = r
    return used


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def read_jgi_fasta(path: Path, portal: str, wanted: set[str]) -> dict[str, tuple[str, str]]:
    """{proteinId: (header name, sequence)} for the wanted IDs of one gzip FASTA.

    The header is `>jgi|<portal>|<proteinId>|<name...>`. A header that does not follow it, or
    names another portal, is skipped. The file is data: it is only read, never executed.
    """
    out: dict[str, tuple[str, str]] = {}
    pid, name, chunks = None, "", []

    def flush():
        if pid is not None:
            seq = "".join(chunks).upper().rstrip("*")
            out[pid] = (name, seq)

    with gzip.open(path, "rt", encoding="utf-8", errors="replace") as fh:
        for line in fh:
            line = line.rstrip("\n")
            if line.startswith(">"):
                flush()
                pid, chunks = None, []
                parts = line[1:].split("|", 3)
                if len(parts) == 4 and parts[0] == "jgi" and parts[1] == portal:
                    if parts[2] in wanted:
                        pid, name = parts[2], parts[3].strip()
            elif pid is not None:
                chunks.append(line.strip())
    flush()
    return out


def load_jgi_proteins(
    jgi_dir: Path, table: dict[str, dict], wanted: dict[str, set[str]]
) -> dict[str, dict[str, tuple[str, str]]]:
    """Verify the sha256 of each used file against the table, then read the wanted IDs."""
    out: dict[str, dict[str, tuple[str, str]]] = {}
    for portal in sorted(wanted):
        if portal not in table:
            continue
        found = sorted((jgi_dir / portal).rglob(table[portal]["file"]))
        if len(found) != 1:
            raise FileNotFoundError(f"{portal}: {len(found)} copies of {table[portal]['file']}")
        got = file_sha256(found[0])
        if got != table[portal]["sha256"]:
            raise ValueError(f"{portal}: sha256 {got} differs from the file table")
        out[portal] = read_jgi_fasta(found[0], portal, wanted[portal])
    return out


def header_variants(name: str) -> set[str]:
    """Upper-case forms of the first header token that a UniProt gene name could carry.

    CAALFM_CR00610WAT0 -> CAALFM_CR00610WA, NCU06242T0 -> NCU06242, BCIN_02g00390.1 ->
    BCIN_02G00390, mRNA_UMAG_11562 -> UMAG_11562.
    """
    tok = name.split()[0].upper() if name.split() else ""
    out = {tok} if tok else set()
    for v in list(out):
        for rx in (r"^MRNA_", r"\.\d+$", r"T\d+$"):
            w = re.sub(rx, "", v)
            if w:
                out.add(w)
    return out


def uniprot_gene_tokens(genes: str) -> set[str]:
    return {t.upper() for t in re.split(r"[\s;,]+", genes.strip()) if t}


def jgi_cross_check(row: dict, jgi_name: str, jgi_seq: str) -> tuple[str, str, str]:
    """(check, how, vs_uniprot) of one found JGI protein against what the row carries.

    check: agree | differs | unverifiable. Evidence in order: a locus tag of the row (only a
    tag of the same naming scheme as the header can agree or differ), a gene name of the row
    equal to the header name, the UniProt candidate (its gene names carry the header locus, or
    its sequence is identical; a similar but not identical sequence is unverifiable, because
    the candidate may come from another strain). `vs_uniprot` is identical | similar |
    different | no_uniprot_sequence, whatever the check.
    """
    uni_seq = row.get("sequence", "")
    if not uni_seq:
        vs = "no_uniprot_sequence"
    elif jgi_seq == uni_seq:
        vs = "identical"
    else:
        short = max(len(jgi_seq), len(uni_seq)) <= SIMILAR_MAX_LEN
        ratio = (
            difflib.SequenceMatcher(None, jgi_seq, uni_seq, autojunk=False).ratio()
            if short
            else 0.0
        )
        vs = "similar" if ratio >= SIMILAR_RATIO else "different"
    variants = header_variants(jgi_name)
    first = jgi_name.split()[0].lower() if jgi_name.split() else ""
    tags = locus_tags(row["gene_or_family_id"])
    same_scheme = False
    for t in tags:
        tu = t.upper()
        if tu in variants or any(v.endswith("_" + tu) for v in variants):
            return "agree", "row_locus_tag", vs
        # UM05433 and UMAG_05433 share the number; CND01240 and CNAG_00996 are not comparable
        letters = re.match(r"[A-Z]+", tu)
        for v in variants:
            vl = re.match(r"[A-Z]+", v.replace("_", ""))
            if letters and vl and vl.group(0).startswith(letters.group(0)):
                if re.sub(r"\D", "", tu) == re.sub(r"\D", "", v):
                    return "agree", "row_locus_tag_same_number", vs
                same_scheme = True
    if tags and same_scheme:
        return "differs", "row_locus_tag_not_in_header", vs
    if tags:
        return "unverifiable", "row_locus_tag_other_naming_scheme", vs
    if first and first in {t.lower() for t in gene_name_tokens(row["gene_or_family_id"])}:
        return "agree", "row_gene_name_equals_header_name", vs
    if uni_seq and row.get("map_status") in ("candidate_gene_name", "mapped_locus_tag"):
        if variants & uniprot_gene_tokens(row.get("genes", "")):
            return "agree", "uniprot_gene_names_carry_header_locus", vs
        if vs == "identical":
            return "agree", "uniprot_sequence_identical", vs
        if vs == "similar":
            return "unverifiable", "similar_not_identical_to_uniprot_candidate", vs
        return "differs", "differs_from_uniprot_candidate", vs
    return "unverifiable", "no_name_or_uniprot_evidence", vs


def apply_jgi(
    rows: list[dict], proteins: dict[str, dict[str, tuple[str, str]]], provided: set[str]
) -> dict:
    """Add the JGI columns and decide mapped_jgi_file. Returns {portal: evidence counts}.

    Rows mapped by locus tag keep their status (the file only cross-checks them). A candidate
    row keeps its status unless the file accepts it (mapped_jgi_file) or contradicts it
    (lost). Every other row without evidence stays lost with the file reason appended.
    """
    for r in rows:
        for c in JGI_COLS:
            r[c] = ""
        r["map_status_uniprot"] = r["map_status"]
    stats: dict[str, dict[str, int]] = defaultdict(
        lambda: {"rows": 0, "found": 0, "absent": 0, "agree": 0, "differs": 0, "unverifiable": 0}
    )
    todo = [r for r in rows if r["organism_code"]]
    for r in todo:
        portal = jgi_portal(r["organism_code"])
        r["jgi_portal"] = portal
        if portal not in provided:
            r["jgi_status"] = "portal_not_provided"
            continue
        st = stats[portal]
        st["rows"] += 1
        hit = proteins.get(portal, {}).get(jgi_protein_id(r["protein_id_jgi"]))
        if hit is None:
            r["jgi_status"] = "absent"
            st["absent"] += 1
            continue
        name, seq = hit
        r["jgi_status"] = "found"
        st["found"] += 1
        r["jgi_name"] = name
        r["jgi_length"] = str(len(seq))
        r["jgi_seq_sha256"] = hashlib.sha256(seq.encode()).hexdigest()
        r["jgi_check"], r["jgi_check_how"], r["jgi_vs_uniprot"] = jgi_cross_check(r, name, seq)
        st[r["jgi_check"]] += 1
    validated = {
        p
        for p, st in stats.items()
        if st["agree"] >= PORTAL_MIN_AGREE
        and st["differs"] <= PORTAL_MAX_DIFFER_RATE * (st["agree"] + st["differs"])
    }
    for r in todo:
        if r["jgi_status"] == "portal_not_provided":
            note = "jgi_portal_not_provided"
        elif r["jgi_status"] == "absent":
            note = "jgi_id_absent_from_file"
        else:
            ok, why = plausible_protein(
                proteins[r["jgi_portal"]][jgi_protein_id(r["protein_id_jgi"])][1]
            )
            if not ok:
                note = f"jgi_sequence_implausible: {why}"
            elif r["jgi_check"] == "agree":
                r["jgi_accept"] = "row_verified"
            elif r["jgi_check"] == "unverifiable" and r["jgi_portal"] in validated:
                r["jgi_accept"] = "portal_validated"
            else:
                note = (
                    "jgi_file_" + r["jgi_check_how"]
                    if r["jgi_check"] == "differs"
                    else "jgi_file_unverifiable_portal_not_validated"
                )
        if r["jgi_accept"] and r["map_status"] != "mapped_locus_tag":
            r["map_status"], r["map_reason"] = "mapped_jgi_file", ""
            continue
        if r["jgi_accept"]:
            continue
        if r["map_status"] == "candidate_gene_name" and r["jgi_check"] != "differs":
            continue  # still only a candidate; the file did not decide
        if r["map_status"] == "mapped_locus_tag":
            continue
        r["map_status"] = "lost"
        r["map_reason"] = f"{r['map_reason']}; {note}" if r["map_reason"] else note
    for p in stats:
        stats[p]["portal_validated"] = int(p in validated)
    return dict(stats)


def header_locus_tag(name: str) -> str:
    """The locus tag of a JGI header name, or '' when the name is a gene-model name.

    HCAG_00594T0 -> HCAG_00594, BCIN_02g00390.1 -> BCIN_02g00390, mRNA_UMAG_05433 -> UMAG_05433.
    fgenesh1_kg.2_#_1532_#_Locus43v1rpkm3058.88 and Crypa1.estExt_fgenesh1_kg.C_90297 give ''.
    """
    tok = name.split()[0] if name.split() else ""
    tok = re.sub(r"^mRNA_", "", tok)
    tok = re.sub(r"\.\d+$", "", tok)
    if not LOCUS_LIKE.match(tok):
        return ""
    stripped = re.sub(r"T\d+$", "", tok)
    return stripped if LOCUS_LIKE.match(stripped) else tok


def apply_locus_lookup(rows: list[dict], proteins: dict, scope_of, lookup) -> dict:
    """UniProt lookup by the locus tag of the JGI header, for rows the file step left unmapped.

    Only rows with a found JGI protein and a header locus tag are looked up. An exact single
    hit (reviewed first) whose sequence equals the JGI protein confirms the row (status
    mapped_jgi_file, jgi_accept locus_confirmed). A hit with another sequence is reported
    (`sequence_differs`) and the row stays unmapped. No hit or an ambiguous hit leaves it so.
    The hit proves that the file protein is that UniProt entry. It does not prove that the Kiss
    row means it, so `locus_confirmed` is weaker than `row_verified`. When the row has a
    gene-name candidate that is a different entry (`confirmed_name_conflict`), the row's name
    contradicts the protein at this ID and the row is not mapped.
    """
    stats: dict[str, int] = defaultdict(int)
    for r in rows:
        if (
            not r["organism_code"]
            or r["jgi_status"] != "found"
            or r["jgi_accept"]
            or r["map_status"] == "mapped_locus_tag"
        ):
            continue
        name, seq = proteins[r["jgi_portal"]][jgi_protein_id(r["protein_id_jgi"])]
        tag = header_locus_tag(name)
        if not tag:
            r["jgi_locus_result"] = "no_locus_tag_in_header"
            stats["no_locus_tag_in_header"] += 1
            continue
        r["jgi_locus_tag"] = tag
        entry, why = pick_entry(lookup("gene_exact", tag, scope_of(r["organism_code"])))
        if entry is not None:
            r["jgi_locus_genes"] = entry.get("genes", "")
        if entry is None:
            r["jgi_locus_result"] = why
        elif entry["sequence"].upper() != seq:
            r["jgi_locus_result"], r["jgi_locus_accession"] = "sequence_differs", entry["accession"]
            r["map_reason"] = (
                f"{r['map_reason']}; jgi_header_locus_tag_hit_sequence_differs"
                if r["map_reason"]
                else "jgi_header_locus_tag_hit_sequence_differs"
            )
            if r["map_status"] == "candidate_gene_name":
                r["map_status"] = "lost"
        elif (
            r["map_status_uniprot"] == "candidate_gene_name"
            and r["accession"] != entry["accession"]
        ):
            # the protein at this ID is another gene than the one the row's name points to
            r["jgi_locus_result"], r["jgi_locus_accession"] = (
                "confirmed_name_conflict",
                entry["accession"],
            )
            r["map_reason"] = "jgi_header_locus_tag_entry_differs_from_name_candidate"
            r["map_status"] = "lost"
        else:
            r["jgi_locus_result"], r["jgi_locus_accession"] = "confirmed", entry["accession"]
            r["jgi_accept"] = "locus_confirmed"
            r["map_status"], r["map_reason"] = "mapped_jgi_file", ""
            r["accession"], r["length"] = entry["accession"], str(len(seq))
            r["uniprot_organism_id"] = str(entry["organism_id"])
            r["match_kind"], r["match_term"] = "jgi_header_locus_tag", tag
        stats[r["jgi_locus_result"]] += 1
    return dict(stats)


def locus_scope(code: str, org: dict, cache) -> str:
    """UniProt taxon for the header-locus-tag lookup of one organism.

    The genome strain of Kiss Data 1 (STRAIN_TERMS) when NCBI Taxonomy resolves that exact term
    to one taxon, even when the taxon is now named differently (Histoplasma capsulatum NAm1
    resolves to 2059318, Histoplasma mississippiense). Otherwise the scope of the first stage.
    """
    term = STRAIN_TERMS.get(code)
    if term:
        ids = ncbi_taxid(cache, term)
        if len(ids) == 1:
            return ids[0]
    return org[code]["scope_taxid"]


def guard_no_sequences(paths: list[Path], sequences: list[str]) -> None:
    """Fail when any written table carries a JGI sequence (they stay out of the repository)."""
    seqs = [s for s in sequences if len(s) >= MIN_LEN]
    for p in paths:
        text = p.read_text()
        for s in seqs:
            if s in text:
                raise RuntimeError(f"a JGI sequence of length {len(s)} was written to {p}")


# ----------------------------------------------------------------------------- network + cache


class Cache:
    """URL -> text cache in one JSON file with sorted keys (deterministic bytes)."""

    def __init__(self, path: Path, online: bool):
        self.path, self.online = path, online
        self.data: dict[str, str] = json.loads(path.read_text()) if path.exists() else {}
        self.dirty = False
        self.log: list[str] = []

    def get(self, url: str) -> str:
        if url in self.data:
            return self.data[url]
        if not self.online:
            raise KeyError(f"not in cache and --online not set: {url}")
        req = urllib.request.Request(url, headers={"User-Agent": "adhesionPred-kiss-feasibility"})
        for attempt in range(4):
            try:
                with urllib.request.urlopen(req, timeout=60) as resp:  # noqa: S310 (https only)
                    text = resp.read().decode("utf-8")
                break
            except (urllib.error.URLError, TimeoutError):
                if attempt == 3:
                    raise
                time.sleep(2 * (attempt + 1))
        self.data[url] = text
        self.dirty = True
        time.sleep(0.4)
        return text

    def save(self) -> None:
        if self.dirty:
            self.path.write_text(json.dumps(self.data, indent=0, sort_keys=True) + "\n")


def ncbi_taxid(cache: Cache, term: str) -> list[str]:
    q = urllib.parse.urlencode({"db": "taxonomy", "term": term, "retmode": "json", "retmax": "5"})
    js = json.loads(cache.get(f"{NCBI}esearch.fcgi?{q}"))
    return js["esearchresult"]["idlist"]


def ncbi_taxon(cache: Cache, taxid: str) -> dict:
    q = urllib.parse.urlencode({"db": "taxonomy", "id": taxid, "retmode": "xml"})
    root = ElementTree.fromstring(cache.get(f"{NCBI}efetch.fcgi?{q}"))  # noqa: S314 (public NCBI XML)
    t = root.find("Taxon")
    lineage = {}
    for x in t.findall("LineageEx/Taxon"):
        rank = x.findtext("Rank")
        if rank and rank != "no rank":
            lineage[rank] = (x.findtext("ScientificName"), x.findtext("TaxId"))
    return {
        "taxid": t.findtext("TaxId"),
        "name": t.findtext("ScientificName"),
        "rank": t.findtext("Rank"),
        "lineage": lineage,
    }


def clades_from_lineage(lineage: dict) -> tuple[str, str]:
    """(coarse, fine) clade from the NCBI lineage (never from node_of_emergence).

    coarse: phylum, but for Ascomycota the subphylum.
    fine:   as coarse, but Pezizomycotina and Basidiomycota use the class.
    """
    phylum = lineage.get("phylum", ("", ""))[0]
    sub = lineage.get("subphylum", ("", ""))[0]
    cls = lineage.get("class", ("", ""))[0]
    if phylum == "Ascomycota":
        coarse = sub
        fine = cls if sub == "Pezizomycotina" else sub
    elif phylum == "Basidiomycota":
        coarse, fine = phylum, cls
    else:
        coarse = fine = phylum
    return coarse, fine


def parse_uniprot_tsv(text: str) -> list[dict]:
    lines = text.strip("\n").split("\n")
    out = []
    for ln in lines[1:]:
        f = ln.split("\t")
        if len(f) < 8:
            continue
        out.append(
            {
                "accession": f[0],
                "reviewed": f[1].strip() == "reviewed",
                "entry_name": f[2],
                "genes": f[3],
                "organism": f[4],
                "organism_id": int(f[5]),
                "length": int(f[6]),
                "sequence": f[7],
            }
        )
    return out


def uniprot_query_url(query: str, size: int = 50) -> str:
    q = urllib.parse.urlencode(
        {"query": query, "format": "tsv", "fields": UNIPROT_FIELDS, "size": str(size)}
    )
    return f"{UNIPROT}?{q}"


def make_lookup(cache: Cache):
    def lookup(kind: str, term: str, taxon: str) -> list[dict]:
        query = f"({kind}:{term}) AND (taxonomy_id:{taxon})"
        return parse_uniprot_tsv(cache.get(uniprot_query_url(query)))

    return lookup


# ----------------------------------------------------------------------------- stages


def read_tsv(path: Path) -> list[dict]:
    with path.open(newline="") as fh:
        return list(csv.DictReader(fh, delimiter="\t"))


def write_tsv(path: Path, header: list[str], rows: list[list]) -> None:
    with path.open("w", newline="") as fh:
        w = csv.writer(fh, delimiter="\t", lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)


def species_from_xlsx(xlsx: Path, out: Path) -> None:
    import openpyxl  # noqa: PLC0415 (optional, manual step)

    ws = openpyxl.load_workbook(xlsx, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hdr = ["species_id", "species_name", "complexity", "genome_version", "assembly_id", "source"]
    body = []
    for r in rows[3:]:
        if r[0]:
            clean = [
                str(c or "").replace("\n", " ").strip()
                for c in (r[0], r[1], r[2], r[3], r[5], r[6])
            ]
            body.append(clean)
    write_tsv(out, hdr, sorted(body))


def data3_check(xlsx: Path, rows: list[dict], out: Path) -> None:
    """Re-derive category counts from the full Data 3 sheet (manual step, needs openpyxl)."""
    import openpyxl  # noqa: PLC0415 (optional, manual step)

    ws = openpyxl.load_workbook(xlsx, read_only=True)["gene list"]
    data = [r for r in ws.iter_rows(min_row=4, values_only=True) if r[0] and r[1]]
    # columns: category, gene, JGI id, function, Afu ortholog, cluster, node, source
    tsv_d3 = {
        (r["gene_or_family_id"], r["protein_id_jgi"].strip(), r["cluster_number"])
        for r in rows
        if r["source_sheet"].startswith("Supp Data 3")
    }
    xl = {(str(r[1]).strip(), str(r[2] or "").strip(), str(r[5]).strip()) for r in data}
    per = defaultdict(lambda: {"rows": 0, "clusters": set(), "with_id": 0})
    for r in data:
        d = per[str(r[0]).strip()]
        d["rows"] += 1
        d["clusters"].add(str(r[5]).strip())
        d["with_id"] += 1 if str(r[2] or "").strip() else 0
    body = [[cat, d["rows"], len(d["clusters"]), d["with_id"]] for cat, d in sorted(per.items())]
    body.append(
        [
            "TOTAL",
            len(data),
            len({str(r[5]).strip() for r in data}),
            sum(1 for r in data if str(r[2] or "").strip()),
        ]
    )
    body.append(
        [
            "TSV Data 3 rows found in the xlsx (gene, JGI id, cluster)",
            len(tsv_d3 & xl),
            len(tsv_d3),
            "",
        ]
    )
    write_tsv(
        out / "data3_full_counts.tsv",
        ["functional_category", "rows", "clusters", "with_jgi_id"],
        body,
    )


def resolve_organisms(codes: list[str], species: dict, cache: Cache) -> list[dict]:
    out = []
    for code in sorted(codes):
        sp = species[code]
        name = sp["species_name"]
        ids = ncbi_taxid(cache, f'"{name}"[Scientific Name]')
        how = "species name, Scientific Name"
        if len(ids) != 1:
            ids = ncbi_taxid(cache, f'"{name}"[All Names]')
            how = "species name, All Names"
        if len(ids) != 1:
            out.append({"code": code, "species_name": name, "error": f"NCBI hits {ids}"})
            continue
        sp_tax = ncbi_taxon(cache, ids[0])
        scope, scope_name, scope_how = sp_tax["taxid"], sp_tax["name"], "species taxon"
        term = STRAIN_TERMS.get(code)
        if term:
            sids = ncbi_taxid(cache, f'"{term}"[Scientific Name]')
            if len(sids) == 1:
                st = ncbi_taxon(cache, sids[0])
                if st["name"].lower() == term.lower():
                    scope, scope_name, scope_how = st["taxid"], st["name"], "strain named in Data 1"
        coarse, fine = clades_from_lineage(sp_tax["lineage"])
        lin = sp_tax["lineage"]
        out.append(
            {
                "code": code,
                "species_name_kiss": name,
                "genome_version": sp["genome_version"],
                "ncbi_taxid": sp_tax["taxid"],
                "ncbi_name": sp_tax["name"],
                "name_check": "same" if sp_tax["name"].lower() == name.lower() else "differs",
                "resolve_how": how,
                "scope_taxid": scope,
                "scope_name": scope_name,
                "scope_how": scope_how,
                "phylum": lin.get("phylum", ("", ""))[0],
                "subphylum": lin.get("subphylum", ("", ""))[0],
                "class": lin.get("class", ("", ""))[0],
                "clade_coarse": coarse,
                "clade_fine": fine,
                "clade_taxids": ";".join(
                    f"{r}:{lin[r][1]}" for r in ("phylum", "subphylum", "class") if r in lin
                ),
                "source": "Kiss 2019 Supp Data 1 (MOESM5) species table; NCBI Taxonomy efetch",
            }
        )
    return out


def second_source(models: list[dict], clades: dict, cache: Cache) -> tuple[list[list], list[list]]:
    """Reviewed UniProt entries per wall/fusion Pfam model and clade (counts only).

    Returns (rows per model and clade, summary per model set and clade with distinct entries).
    """
    out = []
    pooled: dict = defaultdict(lambda: (set(), set()))
    for m in sorted(models, key=lambda x: x["id"]):
        if not (m["set"].startswith("wall") or m["set"].startswith("cell_fusion")):
            continue
        for cname, taxid in sorted(clades.items()):
            q = f"(xref:pfam-{m['id']}) AND (reviewed:true) AND (taxonomy_id:{taxid})"
            hits = parse_uniprot_tsv(cache.get(uniprot_query_url(q, size=500)))
            orgs = sorted({h["organism_id"] for h in hits})
            acc_set, org_set = pooled[(m["set"], cname)]
            acc_set.update(h["accession"] for h in hits)
            org_set.update(orgs)
            out.append(
                [
                    m["id"],
                    m["name"],
                    m["set"],
                    cname,
                    taxid,
                    len(hits),
                    len(orgs),
                    len(hits) >= 500,
                    q,
                ]
            )
    summary = [[sset, cname, len(a), len(o)] for (sset, cname), (a, o) in sorted(pooled.items())]
    return out, summary


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--tsv", type=Path, default=TSV_DEFAULT)
    ap.add_argument("--models", type=Path, default=MODELS_DEFAULT)
    ap.add_argument("--out-dir", type=Path, default=OUT_DEFAULT)
    ap.add_argument("--online", action="store_true", help="query public NCBI and UniProt APIs")
    ap.add_argument("--data3-xlsx", type=Path, help="Kiss MOESM7 xlsx: re-derive the Data 3 counts")
    ap.add_argument("--jgi-dir", type=Path, help="folder of the downloaded MycoCosm protein files")
    ap.add_argument("--jgi-table", type=Path, help="file table (default: <jgi-dir>/FILE_TABLE.tsv)")
    ap.add_argument("--species-xlsx", type=Path, help="Kiss MOESM5 xlsx: rebuild the species table")
    args = ap.parse_args(argv)

    out = args.out_dir
    out.mkdir(parents=True, exist_ok=True)
    species_tsv = out / "kiss_species_72.tsv"
    if args.species_xlsx:
        species_from_xlsx(args.species_xlsx, species_tsv)
    species = {r["species_id"]: r for r in read_tsv(species_tsv)}
    cache = Cache(out / "query_cache.json", args.online)

    rows = read_tsv(args.tsv)
    for r in rows:
        r["slice"] = slice_of(r["category"]) or ""
        r["organism_code"] = organism_code(r["protein_id_jgi"])
    codes = sorted({r["organism_code"] for r in rows if r["organism_code"]})
    orgs = resolve_organisms(codes, species, cache)
    bad = [o for o in orgs if "error" in o]
    if bad:
        print("organism resolution failed:", bad, file=sys.stderr)
        return 2
    org = {o["code"]: o for o in orgs}
    afu_scope = org["Aspfu1"]["scope_taxid"]

    lookup = make_lookup(cache)
    for r in rows:
        code = r["organism_code"]
        scope = org[code]["scope_taxid"] if code else afu_scope
        m = map_row(r, scope, lookup)
        r["map_status"], r["map_reason"] = m["status"], m["reason"]
        if m["status"] == "lost" and r["protein_id_jgi"].strip():
            r["map_reason"] = jgi_loss_reason(r, m)
        for k in (
            "accession",
            "match_kind",
            "match_term",
            "length",
            "uniprot_organism_id",
            "id_consistent",
        ):
            r[k] = m[k]
        r["sequence"] = m["sequence"]
        r["clade_coarse"] = org[code]["clade_coarse"] if code else ""
        r["clade_fine"] = org[code]["clade_fine"] if code else ""
        r["genes"] = m["genes"]
    cache.save()
    jgi_on = args.jgi_dir is not None
    jgi_stats: dict = {}
    jgi_seqs: list[str] = []
    if jgi_on:
        table_path = args.jgi_table or args.jgi_dir / "FILE_TABLE.tsv"
        table = read_jgi_table(table_path)
        wanted: dict[str, set[str]] = defaultdict(set)
        for r in rows:
            if r["organism_code"]:
                wanted[jgi_portal(r["organism_code"])].add(jgi_protein_id(r["protein_id_jgi"]))
        proteins = load_jgi_proteins(args.jgi_dir, table, wanted)
        jgi_stats = apply_jgi(rows, proteins, set(table))
        scopes = {c: locus_scope(c, org, cache) for c in sorted(org)}
        lookup_stats = apply_locus_lookup(rows, proteins, scopes.__getitem__, lookup)
        print("locus-tag lookup:", dict(sorted(lookup_stats.items())))
        print("locus-tag scopes:", {c: v for c, v in scopes.items() if v != org[c]["scope_taxid"]})
        cache.save()
        jgi_seqs = [seq for per in proteins.values() for _, seq in per.values()]
        write_tsv(
            out / "jgi_files.tsv",
            list(read_tsv(table_path)[0].keys()),
            [list(r.values()) for r in read_tsv(table_path)],
        )
    if args.data3_xlsx:
        data3_check(args.data3_xlsx, rows, out)

    # ---- write tables
    write_tsv(
        out / "organism_resolution.tsv",
        [
            "code",
            "species_name_kiss",
            "genome_version",
            "ncbi_taxid",
            "ncbi_name",
            "name_check",
            "scope_taxid",
            "scope_name",
            "scope_how",
            "phylum",
            "subphylum",
            "class",
            "clade_coarse",
            "clade_fine",
            "clade_taxids",
            "source",
        ],
        [
            [
                o[k]
                for k in (
                    "code",
                    "species_name_kiss",
                    "genome_version",
                    "ncbi_taxid",
                    "ncbi_name",
                    "name_check",
                    "scope_taxid",
                    "scope_name",
                    "scope_how",
                    "phylum",
                    "subphylum",
                    "class",
                    "clade_coarse",
                    "clade_fine",
                    "clade_taxids",
                    "source",
                )
            ]
            for o in orgs
        ],
    )
    for r in rows:
        r["control_role"] = control_role(r)
    row_cols = [
        "gene_or_family_id",
        "cluster_number",
        "category",
        "slice",
        "protein_id_jgi",
        "organism_code",
        "clade_coarse",
        "clade_fine",
        "map_status",
        "map_reason",
        "match_kind",
        "match_term",
        "accession",
        "length",
        "uniprot_organism_id",
        "id_consistent",
        "control_role",
        *(JGI_COLS if jgi_on else []),
    ]
    ordered = sorted(
        rows,
        key=lambda r: (
            r["slice"],
            int(r["cluster_number"]),
            r["organism_code"],
            r["gene_or_family_id"],
            r["protein_id_jgi"],
        ),
    )
    write_tsv(out / "row_mapping.tsv", row_cols, [[r[c] for c in row_cols] for r in ordered])

    # counts per slice: with/without id, mapped, lost
    stat_rows = []
    for sl in (*SLICES, "all"):
        sub = [r for r in rows if sl == "all" or r["slice"] == sl]
        c = count_rows(sub)
        stat_rows.append(
            [
                sl,
                c["rows"],
                len({r["cluster_number"] for r in sub}),
                c["with_jgi_id"],
                c["without_jgi_id"],
                c["mapped_locus_tag"],
                c["distinct_accessions_locus_tag"],
                c["mapped_family_rep"],
                c["candidate_gene_name"],
                c["lost"],
                *([c["mapped_jgi_file"]] if jgi_on else []),
            ]
        )
    write_tsv(
        out / "counts_per_slice.tsv",
        [
            "slice",
            "rows",
            "clusters_kiss",
            "with_jgi_id",
            "without_jgi_id",
            "mapped_locus_tag",
            "distinct_accessions_locus_tag",
            "mapped_family_rep",
            "candidate_gene_name",
            "lost",
            *(["mapped_jgi_file"] if jgi_on else []),
        ],
        stat_rows,
    )
    reasons = defaultdict(int)
    for r in rows:
        if r["map_status"] == "lost":
            reasons[(r["slice"], r["map_reason"])] += 1
    write_tsv(
        out / "loss_reasons.tsv",
        ["slice", "reason", "rows"],
        [[s, why, n] for (s, why), n in sorted(reasons.items())],
    )
    per_org = defaultdict(lambda: defaultdict(int))
    for r in rows:
        per_org[r["organism_code"] or "(no JGI ID)"][(r["slice"], r["map_status"])] += 1
    po_rows = []
    for code in sorted(per_org):
        d = per_org[code]
        for sl in SLICES:
            tot = sum(v for (s, _), v in d.items() if s == sl)
            if tot:
                po_rows.append(
                    [
                        code,
                        org[code]["clade_coarse"] if code in org else "",
                        sl,
                        tot,
                        d[(sl, "mapped_locus_tag")],
                        d[(sl, "mapped_family_rep")],
                        d[(sl, "candidate_gene_name")],
                        d[(sl, "lost")],
                        *([d[(sl, "mapped_jgi_file")]] if jgi_on else []),
                    ]
                )
    write_tsv(
        out / "rows_per_organism.tsv",
        [
            "organism_code",
            "clade_coarse",
            "slice",
            "rows",
            "mapped_locus_tag",
            "mapped_family_rep",
            "candidate_gene_name",
            "lost",
            *(["mapped_jgi_file"] if jgi_on else []),
        ],
        po_rows,
    )

    # the first three scenarios are the UniProt-only ones of the first feasibility run
    pre = "map_status_uniprot" if jgi_on else "map_status"
    scenarios = {
        "upper_bound_all_jgi_id_rows": lambda r: bool(r["organism_code"]),
        "strict_locus_tag_mapped": lambda r: r[pre] == "mapped_locus_tag",
        "locus_tag_plus_gene_name_candidates": lambda r: r[pre]
        in ("mapped_locus_tag", "candidate_gene_name"),
    }
    if jgi_on:
        scenarios["strict_plus_jgi_file"] = lambda r: r["map_status"] in (
            "mapped_locus_tag",
            "mapped_jgi_file",
        )
        scenarios["strict_plus_jgi_file_row_verified"] = lambda r: r["map_status"] == (
            "mapped_locus_tag"
        ) or (r["map_status"] == "mapped_jgi_file" and r["jgi_accept"] == "row_verified")
    clade_rows, floor_rows = [], []
    for sname, inc in scenarios.items():
        tabs = clade_tables(rows, inc)
        for (sl, level, clade), v in sorted(tabs.items()):
            only = v["clusters"] - v["shared"]
            clade_rows.append(
                [
                    sname,
                    sl,
                    level,
                    level_role(level),
                    clade,
                    v["rows"],
                    v["clusters"],
                    v["shared"],
                    only,
                    stratum_status(v["clusters"], sl),
                    stratum_status(v["shared"], sl),
                    stratum_status(only, sl),
                ]
            )
        for sl in SLICES:
            for level in ("coarse", "fine"):
                n_clades = len({c for (s, lv, c) in tabs if s == sl and lv == level})
                for floor in FLOORS:
                    el = eligible(tabs, sl, level, floor)
                    floor_rows.append(
                        [
                            sname,
                            sl,
                            level,
                            floor,
                            n_clades,
                            len(el),
                            "; ".join(f"{c} ({k} clusters, {n} rows)" for c, k, n in el),
                            level_role(level),
                        ]
                    )
    write_tsv(
        out / "clusters_per_clade.tsv",
        [
            "scenario",
            "slice",
            "clade_level",
            "role",
            "clade",
            "rows",
            "clusters",
            "clusters_also_in_other_clade",
            "clusters_clade_only",
            "status_all_clusters",
            "status_shared",
            "status_clade_only",
        ],
        clade_rows,
    )
    write_tsv(
        out / "floors_per_slice.tsv",
        [
            "scenario",
            "slice",
            "clade_level",
            "floor",
            "clades_present",
            "clades_eligible",
            "eligible_clades",
            "role",
        ],
        floor_rows,
    )

    # cluster overlap of family-level (no JGI ID) rows with ID-bearing rows
    id_clusters = {(r["slice"], r["cluster_number"]) for r in rows if r["organism_code"]}
    noid = [r for r in rows if not r["organism_code"]]
    new_clusters = sorted({(r["slice"], r["cluster_number"]) for r in noid} - id_clusters)
    write_tsv(
        out / "family_rows_cluster_overlap.tsv",
        ["slice", "cluster_number", "in_id_bearing_rows"],
        [[s, c, "no"] for s, c in new_clusters]
        + [
            [s, c, "yes"]
            for s, c in sorted({(r["slice"], r["cluster_number"]) for r in noid} & id_clusters)
        ],
    )

    write_tsv(
        out / "family_rep_clusters.tsv",
        ["slice", "cluster_number", "family_rep_rows", "strictly_mapped_kiss_rows"],
        family_rep_clusters(rows),
    )
    if jgi_on:
        write_tsv(out / "name_conflict_review.tsv", NAME_CONFLICT_COLS, name_conflict_table(rows))

    # reconciliation with report 06 / paper
    recon = []
    d3 = [r for r in rows if r["source_sheet"].startswith("Supp Data 3")]
    for sl in SLICES:
        a = [r for r in rows if r["slice"] == sl]
        b = [r for r in d3 if r["slice"] == sl]
        recon.append(
            [
                sl,
                len(a),
                len({r["cluster_number"] for r in a}),
                len(b),
                len({r["cluster_number"] for r in b}),
            ]
        )
    write_tsv(
        out / "reconciliation.tsv",
        [
            "slice",
            "rows_all_sheets",
            "clusters_all_sheets",
            "rows_data3_only",
            "clusters_data3_only",
        ],
        recon,
    )

    # second source
    models = read_tsv(args.models)
    taxa = {
        "Basidiomycota": None,
        "Saccharomycotina": None,
        "Chytridiomycota": None,
        "Mucoromycota": None,
    }
    # clade taxids come from the NCBI lineage of the resolved organisms where present
    for o in orgs:
        for part in o["clade_taxids"].split(";"):
            rank, tid = part.split(":")
            name = {"phylum": o["phylum"], "subphylum": o["subphylum"]}.get(rank)
            if name in taxa:
                taxa[name] = tid
    wanted = {k: v for k, v in taxa.items() if v}
    ss, ss_sum = second_source(models, wanted, cache)
    cache.save()
    write_tsv(
        out / "second_source_summary.tsv",
        ["model_set", "clade", "distinct_reviewed_entries", "distinct_organisms"],
        ss_sum,
    )
    write_tsv(
        out / "second_source_uniprot_reviewed.tsv",
        [
            "pfam_id",
            "model_name",
            "model_set",
            "clade",
            "clade_taxid",
            "reviewed_entries",
            "distinct_organisms",
            "hit_cap_reached",
            "query",
        ],
        ss,
    )
    if jgi_on:
        summ = []
        for portal in sorted(jgi_stats):
            st = jgi_stats[portal]
            for sl in SLICES:
                sub = [r for r in rows if r["jgi_portal"] == portal and r["slice"] == sl]
                if not sub:
                    continue

                def n(pred, sub=sub):
                    return sum(1 for r in sub if pred(r))

                summ.append(
                    [
                        portal,
                        sl,
                        len(sub),
                        n(lambda r: r["jgi_status"] == "found"),
                        n(lambda r: r["jgi_status"] == "absent"),
                        n(lambda r: r["jgi_check"] == "agree"),
                        n(lambda r: r["jgi_check"] == "differs"),
                        n(lambda r: r["jgi_check"] == "unverifiable"),
                        st["portal_validated"],
                        n(lambda r: r["jgi_accept"] == "row_verified"),
                        n(lambda r: r["jgi_accept"] == "locus_confirmed"),
                        n(lambda r: r["jgi_accept"] == "portal_validated"),
                        n(lambda r: r["jgi_vs_uniprot"] == "identical"),
                        n(lambda r: r["jgi_vs_uniprot"] == "different"),
                    ]
                )
        write_tsv(
            out / "jgi_portal_summary.tsv",
            [
                "portal",
                "slice",
                "rows",
                "found",
                "absent",
                "check_agree",
                "check_differs",
                "check_unverifiable",
                "portal_validated",
                "accepted_row_verified",
                "accepted_locus_confirmed",
                "accepted_portal_validated",
                "sequence_identical_to_uniprot",
                "sequence_differs_from_uniprot",
            ],
            summ,
        )
        # the cache and mapped_sequences.faa hold public UniProt sequences, so they are not scanned
        guard_no_sequences(sorted(out.glob("*.tsv")), jgi_seqs)
    sha = hashlib.sha256((out / "row_mapping.tsv").read_bytes()).hexdigest()
    print(f"rows {len(rows)}; row_mapping.tsv sha256 {sha}")

    # mapped sequences (small FASTA, derived from UniProt; not JGI sequences)
    fa = []
    for r in ordered:
        if r["map_status"] in ("mapped_locus_tag", "mapped_family_rep") and r["accession"]:
            fa.append((r["accession"], r["sequence"]))
    seen, lines = set(), []
    for acc, seq in sorted(fa):
        if acc not in seen:
            seen.add(acc)
            lines += [f">{acc}", seq]
    if sum(len(x) for x in lines) < 1_000_000:
        (out / "mapped_sequences.faa").write_text("\n".join(lines) + ("\n" if lines else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main())
